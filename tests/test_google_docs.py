"""Bringing in a shared Google Doc as a retreat's source."""

import io
import json
import time

import docx
import httpx
import pytest
from fastapi.testclient import TestClient

from app import google_docs, main, tts
from app.auth import User, current_user

VOICES = {"guide": "en-US-AvaMultilingualNeural", "reading": "en-US-AndrewMultilingualNeural",
          "heart": "en-US-EmmaMultilingualNeural", "deep": "en-US-ChristopherNeural"}
DOC = "https://docs.google.com/document/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit?usp=sharing"


def _docx() -> bytes:
    d = docx.Document()
    d.add_paragraph("Day 1: Be still")
    d.add_paragraph("Psalm 46:10. Be still, and know that I am God.")
    d.add_paragraph("Day 2: My shepherd")
    d.add_paragraph("Psalm 23:1. Yahweh is my shepherd: I shall lack nothing.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _mock(monkeypatch, handler):
    real = httpx.AsyncClient
    monkeypatch.setattr(httpx, "AsyncClient", lambda **kw: real(transport=httpx.MockTransport(handler), **kw))


def test_doc_id_from_links():
    assert google_docs.doc_id(DOC) == "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789"
    assert google_docs.doc_id("https://docs.google.com/document/u/1/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/view")
    with pytest.raises(google_docs.GoogleDocError):
        google_docs.doc_id("https://example.com/not-a-doc")


def test_a_shared_doc_becomes_a_retreat(monkeypatch):
    def handler(request):
        assert request.url.path.endswith("/export") and request.url.params["format"] == "docx"
        return httpx.Response(200, content=_docx(), headers={
            "content-type": google_docs.DOCX_TYPE,
            "content-disposition": "attachment; filename=\"Psalms.docx\"; filename*=UTF-8''Two%20Psalms.docx"})

    async def fake_synthesize(text, voice, out_path):
        out_path.write_bytes(b"ID3")
        return 2.0

    _mock(monkeypatch, handler)
    monkeypatch.setattr(tts, "synthesize", fake_synthesize)
    with TestClient(main.app) as client:
        main.app.dependency_overrides[current_user] = lambda: User("gdoc-user", "g@x.y")
        r = client.post("/api/retreats", data={"google_doc": DOC, "options": json.dumps({"voices": VOICES})})
        assert r.status_code == 202, r.text
        assert r.json()["filename"] == "Two Psalms.docx"
        for _ in range(200):
            if client.get(f"/api/retreats/{r.json()['id']}").json()["status"] == "ready":
                break
            time.sleep(0.05)
        main.app.dependency_overrides.clear()


def test_a_private_doc_says_how_to_share_it(monkeypatch):
    _mock(monkeypatch, lambda request: httpx.Response(200, text="<html>Sign in</html>", headers={"content-type": "text/html"}))
    with TestClient(main.app) as client:
        main.app.dependency_overrides[current_user] = lambda: User("gdoc-user-2", "g@x.y")
        r = client.post("/api/retreats", data={"google_doc": DOC})
        main.app.dependency_overrides.clear()
    assert r.status_code == 400 and "Anyone with the link" in r.json()["error"]["message"]
