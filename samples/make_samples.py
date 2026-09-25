"""Build the public-domain sample uploads used in the demo and tests.

Scripture: World English Bible (public domain), via bible-api.com.
Images: Rembrandt, The Return of the Prodigal Son (c. 1668), and Henry Ossawa
Tanner, The Annunciation (1898), both public domain, via Wikimedia Commons.

Run: .venv/bin/python samples/make_samples.py
"""

import json
import re
import urllib.request
from pathlib import Path

import docx
import pymupdf
from docx.shared import Inches

HERE = Path(__file__).parent
SRC = HERE / "src"
STRAIGHT = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"'})
UA = {"User-Agent": "IgnatiusAtHomeHW/0.1 (student project)"}

IMAGES = {
    "prodigal.jpg": "https://upload.wikimedia.org/wikipedia/commons/thumb/9/93/Rembrandt_Harmensz_van_Rijn_-_Return_of_the_Prodigal_Son_-_Google_Art_Project.jpg/960px-Rembrandt_Harmensz_van_Rijn_-_Return_of_the_Prodigal_Son_-_Google_Art_Project.jpg",
    "annunciation.jpg": "https://upload.wikimedia.org/wikipedia/commons/thumb/a/ac/Henry_Ossawa_Tanner_-_The_Annunciation.jpg/960px-Henry_Ossawa_Tanner_-_The_Annunciation.jpg",
}


def fetch(url: str) -> bytes:
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=30) as response:
        return response.read()


def image(name: str) -> Path:
    path = SRC / name
    if not path.exists():
        path.write_bytes(fetch(IMAGES[name]))
    return path


def passage(ref: str) -> str:
    path = SRC / (re.sub(r"\W+", "_", ref).strip("_") + ".txt")
    if not path.exists():
        data = json.loads(fetch(f"https://bible-api.com/{ref.replace(' ', '+')}?translation=web"))
        path.write_text(" ".join(v["text"].strip() for v in data["verses"]))
    return path.read_text()


def loose_pdf() -> None:
    """Loose material: three passages and one painting, no days. The app composes a retreat."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 50
    blocks = [("Passages for prayer: rest and return", 16), ("World English Bible (public domain)", 10)]
    for ref in ("Psalm 23", "Matthew 11:28-30", "Luke 15:11-24"):
        blocks += [(ref, 13), (passage(ref), 11)]
    for text, size in blocks:
        text = text.translate(STRAIGHT)  # the built-in PDF fonts lack curly quotes
        rect = pymupdf.Rect(50, y, 545, 800)
        spare = page.insert_textbox(rect, text, fontsize=size, fontname="helv" if size < 13 else "hebo")
        if spare < 0:  # didn't fit; continue on a new page
            page = doc.new_page()
            y = 50
            spare = page.insert_textbox(pymupdf.Rect(50, y, 545, 800), text, fontsize=size, fontname="helv")
        y = 800 - spare + 12
    page = doc.new_page()
    page.insert_text((50, 60), "Rembrandt, The Return of the Prodigal Son (public domain)", fontsize=11)
    page.insert_image(pymupdf.Rect(50, 80, 545, 780), filename=str(image("prodigal.jpg")))
    doc.save(HERE / "loose-passages-web.pdf")


def structured_docx() -> None:
    """Material that already has days. The app keeps them as they are."""
    d = docx.Document()
    d.add_heading("A Three-Day Retreat: Called by Name", level=1)
    d.add_paragraph("Grace to ask for: to hear God call me by name and to answer freely.")
    d.add_heading("Day 1: Isaiah 43:1-4", level=2)
    d.add_paragraph(passage("Isaiah 43:1-4"))
    d.add_heading("Day 2: Luke 1:26-38", level=2)
    d.add_paragraph(passage("Luke 1:26-38"))
    d.add_picture(str(image("annunciation.jpg")), width=Inches(4.5))
    d.add_paragraph("Henry Ossawa Tanner, The Annunciation, 1898 (public domain).")
    d.add_heading("Day 3: Review and savor", level=2)
    d.add_paragraph(
        "Return to the passage from Day 1 or Day 2 where you felt the most consolation or desolation. "
        "Read it slowly again and stay with the word or phrase that stirred you."
    )
    d.save(HERE / "three-days-called-by-name-web.docx")


if __name__ == "__main__":
    SRC.mkdir(exist_ok=True)
    loose_pdf()
    structured_docx()
    print("wrote", *sorted(p.name for p in HERE.glob("*.pdf")), *sorted(p.name for p in HERE.glob("*.docx")))
